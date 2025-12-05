"""
Utility functions for text processing and file handling.
"""
import re
from PyPDF2 import PdfReader
from typing import Union, BinaryIO

from app.logger import setup_logger

logger = setup_logger(__name__)

# Try to import docx support
try:
    from docx import Document
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False
    logger.warning("python-docx not installed. DOCX support disabled.")


def clean_text(text: str) -> str:
    """
    Clean and normalize text by removing HTML tags, URLs, and special characters.

    Args:
        text: Raw text to clean

    Returns:
        Cleaned text with normalized whitespace

    Example:
        >>> clean_text("<p>Hello   World!</p>")
        'Hello World'
    """
    if not text:
        return ""

    # Remove HTML tags
    text = re.sub(r'<[^>]*?>', '', text)

    # Remove URLs
    text = re.sub(
        r'http[s]?://(?:[a-zA-Z]|[0-9]|[$-_@.&+]|[!*\\(\\),]|(?:%[0-9a-fA-F][0-9a-fA-F]))+',
        '',
        text
    )

    # Remove special characters (keep alphanumeric and spaces)
    text = re.sub(r'[^a-zA-Z0-9 ]', ' ', text)

    # Replace multiple spaces with a single space
    text = re.sub(r'\s{2,}', ' ', text)

    # Trim and normalize whitespace
    text = ' '.join(text.split())

    return text


def extract_text_from_pdf(uploaded_file) -> str:
    """
    Extract text content from an uploaded PDF file.

    Args:
        uploaded_file: Streamlit UploadedFile object or file-like object

    Returns:
        Extracted text from all pages

    Raises:
        ValueError: If PDF is empty or cannot be read
        Exception: If PDF extraction fails
    """
    try:
        logger.info("Extracting text from PDF")

        # Read the PDF
        reader = PdfReader(uploaded_file)

        if len(reader.pages) == 0:
            raise ValueError("PDF file is empty")

        # Extract text from all pages
        text_parts = []
        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            else:
                logger.warning(f"No text extracted from page {page_num}")

        full_text = " ".join(text_parts)

        if not full_text.strip():
            raise ValueError("No text could be extracted from the PDF")

        logger.info(f"Successfully extracted {len(full_text)} characters from PDF")
        return full_text

    except Exception as e:
        logger.error(f"Failed to extract text from PDF: {str(e)}")
        raise


def extract_text_from_docx(uploaded_file) -> str:
    """
    Extract text content from an uploaded DOCX file.

    Args:
        uploaded_file: Streamlit UploadedFile object or file-like object

    Returns:
        Extracted text from document

    Raises:
        ImportError: If python-docx is not installed
        ValueError: If DOCX is empty or cannot be read
        Exception: If DOCX extraction fails
    """
    if not DOCX_SUPPORT:
        raise ImportError(
            "python-docx is not installed. Install it with: pip install python-docx"
        )

    try:
        logger.info("Extracting text from DOCX")

        # Read the DOCX
        doc = Document(uploaded_file)

        if len(doc.paragraphs) == 0:
            raise ValueError("DOCX file appears to be empty")

        # Extract text from all paragraphs
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        full_text = " ".join(text_parts)

        if not full_text.strip():
            raise ValueError("No text could be extracted from the DOCX")

        logger.info(f"Successfully extracted {len(full_text)} characters from DOCX")
        return full_text

    except Exception as e:
        logger.error(f"Failed to extract text from DOCX: {str(e)}")
        raise


def extract_text_from_resume(uploaded_file, file_type: str = "pdf") -> str:
    """
    Extract text from resume file (PDF or DOCX).

    Args:
        uploaded_file: Streamlit UploadedFile object
        file_type: File extension ('pdf' or 'docx')

    Returns:
        Extracted text from resume

    Raises:
        ValueError: If file type is not supported
    """
    file_type = file_type.lower().strip()

    if file_type == "pdf":
        return extract_text_from_pdf(uploaded_file)
    elif file_type in ["docx", "doc"]:
        return extract_text_from_docx(uploaded_file)
    else:
        raise ValueError(f"Unsupported file type: {file_type}. Supported: pdf, docx")


def truncate_text(text: str, max_length: int = 1000, suffix: str = "...") -> str:
    """
    Truncate text to a maximum length, adding a suffix if truncated.

    Args:
        text: Text to truncate
        max_length: Maximum length of returned text
        suffix: Suffix to add if text is truncated

    Returns:
        Truncated text with suffix if applicable
    """
    if len(text) <= max_length:
        return text

    return text[:max_length - len(suffix)] + suffix


def format_job_for_display(job: dict) -> str:
    """
    Format a job dictionary for display in the UI.

    Args:
        job: Job dictionary with role, experience, skills, description

    Returns:
        Formatted string representation of the job
    """
    role = job.get('role', 'Unknown Role')
    experience = job.get('experience', 'Not specified')

    skills = job.get('skills', {})
    if isinstance(skills, dict):
        required_skills = skills.get('required', [])
        desired_skills = skills.get('desired', [])
    else:
        required_skills = []
        desired_skills = []

    description = job.get('description', 'No description available')

    formatted = f"**Role:** {role}\n\n"
    formatted += f"**Experience:** {experience}\n\n"

    if required_skills:
        formatted += "**Required Skills:**\n"
        for skill in required_skills:
            formatted += f"- {skill}\n"
        formatted += "\n"

    if desired_skills:
        formatted += "**Desired Skills:**\n"
        for skill in desired_skills:
            formatted += f"- {skill}\n"
        formatted += "\n"

    formatted += f"**Description:**\n{description}"

    return formatted


def validate_url(url: str) -> bool:
    """
    Validate if a string is a valid URL.

    Args:
        url: URL string to validate

    Returns:
        True if valid URL, False otherwise
    """
    url_pattern = re.compile(
        r'^https?://'  # http:// or https://
        r'(?:(?:[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?\.)+[A-Z]{2,6}\.?|'  # domain...
        r'localhost|'  # localhost...
        r'\d{1,3}\.\d{1,3}\.\d{1,3}\.\d{1,3})'  # ...or ip
        r'(?::\d+)?'  # optional port
        r'(?:/?|[/?]\S+)$', re.IGNORECASE
    )
    return url_pattern.match(url) is not None
